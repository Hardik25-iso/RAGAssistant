import io
import os
import tempfile

_tmp = tempfile.mkdtemp()
# CI sets DATABASE_URL to a real Postgres+pgvector service; locally SQLite is used
os.environ.setdefault("DATABASE_URL", f"sqlite:///{_tmp}/test.db")
os.environ.update(
    ENVIRONMENT="test",
    LOCAL_STORAGE_DIR=f"{_tmp}/uploads",
    STORAGE_BACKEND="local",
    EMBEDDINGS_PROVIDER="fake",
    LLM_PROVIDER="fake",
    LLM_FALLBACK_PROVIDER="",
    RATE_LIMIT_ENABLED="false",
    ADMIN_EMAIL="admin@mail.com",
    ADMIN_PASSWORD="AdminPass123",
    JWT_SECRET="test-secret-test-secret-test-secret-123",
)

import fitz  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402

from app.main import app  # noqa: E402


@pytest.fixture(scope="session")
def client():
    with TestClient(app) as c:
        yield c


_counter = {"n": 0}


def register(client, name="Student"):
    _counter["n"] += 1
    email = f"student{_counter['n']}@mail.com"
    r = client.post("/api/auth/register", json={"name": name, "email": email, "password": "Password123"})
    assert r.status_code == 201, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture
def student(client):
    return register(client, "Student A")


@pytest.fixture
def student_b(client):
    return register(client, "Student B")


@pytest.fixture
def admin(client):
    r = client.post("/api/auth/login", json={"email": "admin@mail.com", "password": "AdminPass123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def make_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 800), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_pptx(slides: list[tuple[str, str, str]]) -> bytes:
    """slides: (title, body, speaker_notes)"""
    prs = Presentation()
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def new_subject(client, headers, name="DBMS"):
    r = client.post("/api/subjects", json={"name": name}, headers=headers)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def upload(client, headers, subject_id, filename, data):
    return client.post(
        "/api/documents",
        data={"subject_id": subject_id},
        files={"file": (filename, data, "application/octet-stream")},
        headers=headers,
    )
