"""PPTX -> list of (slide_number, text): title, body, tables, grouped shapes and speaker notes."""
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            out.extend(_shape_text(s))
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
        return out
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        if text.strip():
            out.append(text.strip())
    return out


def load_pptx(data: bytes) -> list[tuple[int, str]]:
    prs = Presentation(io.BytesIO(data))
    slides: list[tuple[int, str]] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        if title:
            parts.append(f"Title: {title}")
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            parts.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append(f"Speaker notes: {notes}")
        text = "\n".join(parts).strip()
        if text:
            slides.append((i, text))
    if not slides:
        raise ValueError("No text found in this presentation")
    return slides
